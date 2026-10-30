"""Owner-controlled uploads with optimistic versions and explicit vector index jobs."""
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import tempfile
import zipfile
import io
from uuid import uuid4
from rag_app.ingest import parse_file
from rag_app.security import AccessError


def parse_upload(filename, content):
    suffix = Path(filename.replace('\\', '/')).suffix.lower()
    if suffix not in ('.txt', '.md', '.pdf', '.docx'):
        raise AccessError('unsupported_file_type', 415)
    if suffix == '.docx':
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as archive:
                entries = archive.infolist()
                if len(entries)>1000 or sum(i.file_size for i in entries)>40*1024*1024:
                    raise AccessError('document_expansion_too_large', 413)
            from docx import Document
            document = Document(io.BytesIO(content))
            parts=[]
            for paragraph in document.paragraphs:
                if not paragraph.text.strip():continue
                style=paragraph.style.name if paragraph.style else ''
                level=style.rsplit(' ',1)[-1]
                prefix=('#'*min(int(level),6)+' ') if style.startswith('Heading ') and level.isdigit() else ''
                parts.append(prefix+paragraph.text)
            for table in document.tables:
                parts.extend(' | '.join(cell.text for cell in row.cells) for row in table.rows)
            text = '\n'.join(parts)
            from rag_app.sections import split_sections
            locations = split_sections(text)
        except AccessError:
            raise
        except Exception:
            raise AccessError('invalid_document', 422) from None
    else:
        try:
            with tempfile.TemporaryDirectory() as folder:
                path = Path(folder) / ('document' + suffix)
                path.write_bytes(content)
                text, locations = parse_file(path)
        except Exception:
            raise AccessError('invalid_or_unextractable_document', 422) from None
    if not text.strip() or not locations:
        raise AccessError('empty_document', 422)
    if len(text)>250_000 or len(locations)>1000:
        raise AccessError('document_text_too_large', 413)
    return text, locations


class Documents:
    def __init__(self, store, state, index=None):
        self.store, self.state, self.index = store, state, index

    def latest(self, actor, doc_id, owner=False):
        with self.store.connect() as c:
            row = self.state.sql(c, 'SELECT * FROM documents WHERE document_id=? ORDER BY version DESC LIMIT 1', (doc_id,)).fetchone()
        if row is None or row['tenant_id'] != actor.tenant_id or (owner and row['owner'] != actor.user_id):
            raise AccessError('document_not_found', 404)
        return dict(row)

    def listing(self, actor):
        rows = {}
        status=self.index.status() if self.index and hasattr(self.index,'status') else None
        for c in self.store.authorized_chunks(actor):
            rows.setdefault(c['document_id'], {k: c[k] for k in ('document_id','title','version','owner','source')})
        return [dict(row, can_manage=row['owner']==actor.user_id,
                     vector_status=status)
                for row in sorted(rows.values(),key=lambda r:r['title'])]

    def upload(self, actor, filename, content, title, groups, doc_id=None, expected_version=None):
        if not isinstance(groups,list) or any(not isinstance(g,str) for g in groups) or len(groups)>20:
            raise AccessError('invalid_document_groups', 422)
        if not set(groups)<=set(actor.groups)|{'public'}:
            raise AccessError('document_group_not_allowed', 403)
        if not title.strip() or len(title)>200:
            raise AccessError('invalid_document_title', 422)
        text, locations = parse_upload(filename,content)
        old = self.latest(actor,doc_id,owner=True) if doc_id else None
        if old and (expected_version is None or expected_version!=old['version']):
            raise AccessError('document_version_conflict', 409)
        clean_name=Path(filename.replace('\\','/')).name[:150]
        doc_id=doc_id or 'upload_'+uuid4().hex
        metadata={'document_id':doc_id,'version':old['version']+1 if old else 1,'title':title.strip(),
            'source':'upload:'+clean_name,'tenant_id':actor.tenant_id,'owner':actor.user_id,
            'allowed_groups':sorted(set(groups)),'created_at':old['created_at'] if old else datetime.now(timezone.utc).isoformat()}
        if old:
            previous_groups=old['allowed_groups'] if isinstance(old['allowed_groups'],list) else json.loads(old['allowed_groups'])
            if (old['active'] and old['hash']==hashlib.sha256(text.encode()).hexdigest() and old['title']==metadata['title']
                and old['source']==metadata['source'] and sorted(previous_groups)==metadata['allowed_groups']):
                return {'document_id':doc_id,'version':old['version'],'changed':False,'vector_index':'unchanged'}
        try:
            self.store.upsert(metadata,text,locations)
        except ValueError:
            raise AccessError('document_version_conflict',409) from None
        return {'document_id':doc_id,'version':metadata['version'],'changed':True,
                'vector_index':'reindex_required' if self.index else 'not_configured', 'lexical_index':'ready'}

    def delete(self, actor, doc_id, expected_version):
        old=self.latest(actor,doc_id,owner=True)
        if old['version']!=expected_version:
            raise AccessError('document_version_conflict',409)
        # Perform ownership, version check and revocation under the store writer lock.
        with self.store.connect() as c:
            if self.store.backend=='postgresql':
                c.execute('SELECT pg_advisory_xact_lock(hashtext(current_schema()),27183)')
            else:
                c.execute('BEGIN IMMEDIATE')
            row=self.state.sql(c,'SELECT version,owner,tenant_id FROM documents WHERE document_id=? ORDER BY version DESC LIMIT 1',(doc_id,)).fetchone()
            if row['version']!=expected_version or row['owner']!=actor.user_id or row['tenant_id']!=actor.tenant_id:
                raise AccessError('document_version_conflict',409)
            self.state.sql(c,'UPDATE documents SET active=0 WHERE document_id=?',(doc_id,))
            c.execute("UPDATE metadata SET value=CAST((CAST(value AS BIGINT)+1) AS TEXT) WHERE key='index_version'")
        return {'document_id':doc_id,'deleted':True,'vector_index':'reindex_required' if self.index else 'not_configured'}

    def queue_index(self, actor, doc_id):
        self.latest(actor,doc_id,owner=True)
        return self.state.create_job(actor,doc_id)

    def run_index_job(self, actor, job_id):
        job=self.state.job(actor,job_id)
        self.latest(actor,job['document_id'],owner=True)
        if not self.state.claim_job(job_id):
            return
        try:
            if self.index:
                if self.store.backend=='postgresql':
                    self.index.sync()
                else:
                    self.index.sync(self.store)
                details=getattr(self.index,'last_sync',{'mode':'full-rebuild'})
            else:
                details={'mode':'lexical-live','index_version':self.store.index_version,'embedded_chunks':0}
            self.state.finish_job(job_id,'succeeded',details)
        except Exception:
            self.state.finish_job(job_id,'failed',{'reason':'index_update_failed','retry':'Submit a new index job after resolving configuration or budgets.'})
