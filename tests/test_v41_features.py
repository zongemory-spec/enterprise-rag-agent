import io
import json
import os
from pathlib import Path
import tempfile
import unittest
from unittest.mock import Mock,patch
import zipfile
import httpx
from fastapi.testclient import TestClient

from init_demo import populate
from rag_app.api import create_app
from rag_app.application_state import ApplicationState
from rag_app.document_management import parse_upload
from rag_app.evidence_policy import EvidencePolicy,HTTPEvidenceJudge
from rag_app.platforms import PlatformRegistry
from rag_app.security import Actor,AccessError,Registry
from rag_app.service import Service
from rag_app.store import Store
from evaluate_feedback import evaluate_records
from export_integrations import specification
from verify_services import run_probes


class NewAPITests(unittest.TestCase):
    def setUp(self):
        self.temp=tempfile.TemporaryDirectory();self.addCleanup(self.temp.cleanup)
        self.root=Path(self.temp.name);self.store=Store(self.root/'db');populate(self.store)
        self.alice=Actor('alice','demo',('support',));self.bob=Actor('bob','demo',('engineering',))
        self.auth=Registry({'alice-test':self.alice,'bob-test':self.bob,'other-test':Actor('alice','other',('support',))})
        entries=[{'platform':p,'token':p+'-test','user_id':'alice','tenant_id':'demo','groups':['support'],
                  'knowledge_ids':['enterprise'],'mode':'service'} for p in ('dify','coze')]
        entries.append({'platform':'coze','token':'delegated-test','user_id':'gateway','tenant_id':'demo',
                        'groups':['support','engineering'],'knowledge_ids':['enterprise'],'mode':'delegated'})
        self.env=patch.dict(os.environ,{'RAG_CHECKPOINT_DB':str(self.root/'cp')})
        self.env.start();self.addCleanup(self.env.stop)
        self.app=create_app(store=self.store,auth=self.auth,platform_auth=PlatformRegistry(entries))
        self.client=TestClient(self.app);self.addCleanup(self.client.close)

    def headers(self,user='alice'):return {'Authorization':'Bearer '+user+'-test'}

    def upload(self,text='The lunar travel reimbursement deadline is 12 days.',**fields):
        data={'title':'Lunar travel reimbursement',**fields}
        return self.client.post('/documents/upload',headers=self.headers(),data=data,
                               files={'file':('policy.md',text.encode(),'text/markdown')})

    def test_upload_is_private_and_searchable(self):
        response=self.upload();self.assertEqual(response.status_code,200,response.text)
        did=response.json()['document_id']
        self.assertTrue(any(d['document_id']==did for d in self.client.get('/documents',headers=self.headers()).json()['documents']))
        self.assertFalse(any(d['document_id']==did for d in self.client.get('/documents',headers=self.headers('bob')).json()['documents']))
        result=self.client.post('/ask',headers=self.headers(),json={'query':'lunar travel reimbursement deadline','method':'bm25'}).json()
        self.assertFalse(result['refused']);self.assertIn('12 days',result['answer'])

    def test_dedup_update_and_old_citations_become_unavailable(self):
        initial=self.upload().json();did=initial['document_id']
        old=next(c for c in self.store.authorized_chunks(self.alice) if c['document_id']==did)
        duplicate=self.upload(document_id=did,expected_version='1')
        self.assertFalse(duplicate.json()['changed'])
        changed=self.upload('The lunar travel reimbursement deadline is 20 days.',document_id=did,expected_version='1')
        self.assertEqual(changed.json()['version'],2)
        self.assertEqual(self.client.get('/citations/'+old['chunk_id'],headers=self.headers()).status_code,404)
        self.assertEqual(self.upload(document_id=did,expected_version='1').status_code,409)

    def test_upload_group_escalation_and_owner_checks(self):
        self.assertEqual(self.upload(groups='["engineering"]').status_code,403)
        created=self.upload(groups='["support"]').json()
        response=self.client.post('/documents/upload',headers=self.headers('bob'),data={
            'title':'Hijack','document_id':created['document_id'],'expected_version':'1'},files={'file':('a.md',b'bad')})
        self.assertEqual(response.status_code,404)
        self.assertEqual(self.upload(groups='not-json').status_code,422)

    def test_upload_size_invalid_format_and_empty_files(self):
        self.assertEqual(self.client.post('/documents/upload',headers=self.headers(),data={'title':'x'},
            files={'file':('bad.exe',b'bad')}).status_code,415)
        self.assertEqual(self.upload('').status_code,422)
        self.assertEqual(self.client.post('/documents/upload',headers=self.headers(),data={'title':'x'},
            files={'file':('big.md',b'x'*(10*1024*1024))}).status_code,413)

    def test_document_delete_invalidates_conversation_and_citations(self):
        doc=self.upload().json()
        answer=self.client.post('/chat',headers=self.headers(),json={'query':'lunar travel reimbursement deadline'}).json()
        cid=answer['conversation_id']
        self.assertEqual(self.client.request('DELETE','/documents/'+doc['document_id'],headers=self.headers(),json={'expected_version':1}).status_code,200)
        turn=self.client.get('/conversations/'+cid,headers=self.headers()).json()['turns'][0]
        self.assertTrue(turn['redacted']);self.assertNotIn('12 days',json.dumps(turn))
        followup=self.client.post('/chat',headers=self.headers(),json={'query':'那需要什么材料？','conversation_id':cid}).json()
        self.assertTrue(followup['refused']);self.assertEqual(followup['refusal_reason'],'followup_requires_clarification')

    def test_index_job_is_explicit_scoped_and_persisted(self):
        document=self.upload(reindex='true').json();jid=document['index_job_id']
        job=self.client.get('/index-jobs/'+jid,headers=self.headers())
        self.assertEqual(job.json()['status'],'succeeded')
        self.assertEqual(job.json()['details']['mode'],'lexical-live')
        self.assertEqual(self.client.get('/index-jobs/'+jid,headers=self.headers('bob')).status_code,404)
        restored=ApplicationState(Store(self.root/'db'))
        self.assertEqual(restored.job(self.alice,jid)['status'],'succeeded')

    def test_followup_is_resolved_using_current_authorized_titles(self):
        first=self.client.post('/chat',headers=self.headers(),json={'query':'差旅报销的规定是什么？'}).json()
        result=self.client.post('/chat',headers=self.headers(),json={'query':'那期限是几天？','conversation_id':first['conversation_id']})
        self.assertEqual(result.status_code,200,result.text)
        self.assertTrue(result.json()['query_rewritten']);self.assertIn('15',result.json()['answer'])
        self.assertEqual(self.client.get('/conversations/'+first['conversation_id'],headers=self.headers('bob')).status_code,404)
        self.assertEqual(self.client.get('/conversations/'+first['conversation_id'],headers=self.headers('other')).status_code,404)

    def test_conversation_cas_and_delete(self):
        state=self.app.state.application_state
        cid=state.create_conversation(self.alice)['conversation_id']
        answer=Service(self.store).ask('差旅报销',self.alice)
        state.append_turn(self.alice,cid,0,'差旅报销',answer)
        with self.assertRaises(AccessError):state.append_turn(self.alice,cid,0,'差旅报销',answer)
        self.assertEqual(self.client.delete('/conversations/'+cid,headers=self.headers()).status_code,200)
        self.assertEqual(self.client.get('/conversations/'+cid,headers=self.headers()).status_code,404)

    def test_feedback_review_export_evaluation_and_revocation(self):
        answer=self.client.post('/ask',headers=self.headers(),json={'query':'退款审批'}).json()
        rid=answer['run_id']
        label={'run_id':rid,'rating':'helpful','reason':'correct','reviewed':True,'expected_refusal':False,'expected_doc_ids':['C001']}
        self.assertEqual(self.client.post('/feedback',headers=self.headers(),json=label).status_code,200)
        self.assertEqual(self.client.post('/feedback',headers=self.headers('bob'),json=label).status_code,403)
        exported=self.client.get('/feedback/export',headers=self.headers()).json()
        self.assertEqual(len(exported['records']),1)
        report=evaluate_records(self.app.state.service,self.alice,exported['records'],['bm25'])
        self.assertEqual(report['results'][0]['evaluated'],1)
        self.assertEqual(report['results'][0]['metrics']['recall_at_1'],1.)
        self.store.deactivate('C001')
        self.assertEqual(self.client.get('/feedback/export',headers=self.headers()).json()['records'],[])

    def test_feedback_no_unreviewed_training_or_inaccessible_labels(self):
        answer=self.client.post('/ask',headers=self.headers(),json={'query':'退款审批'}).json()
        for extra in ({'reviewed':True},{'expected_doc_ids':['E001']}):
            response=self.client.post('/feedback',headers=self.headers(),json={'run_id':answer['run_id'],'rating':'unhelpful',**extra})
            self.assertIn(response.status_code,(404,422))
        self.client.post('/feedback',headers=self.headers(),json={'run_id':answer['run_id'],'rating':'unhelpful'})
        exported=self.client.get('/feedback/export',headers=self.headers()).json()
        self.assertEqual(evaluate_records(self.app.state.service,self.alice,exported['records'],['bm25'])['results'][0]['evaluated'],0)

    def test_dify_external_protocol_and_knowledge_scope(self):
        body={'knowledge_id':'enterprise','query':'退款审批','retrieval_setting':{'top_k':4,'score_threshold':0.0}}
        response=self.client.post('/integrations/dify/retrieval',headers={'Authorization':'Bearer dify-test'},json=body)
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(response.json()['records'])
        row=response.json()['records'][0]
        self.assertEqual(set(row),{'content','score','title','metadata'});self.assertTrue(0<=row['score']<=1)
        self.assertEqual(row['metadata']['document_id'],'C001')
        self.assertEqual(self.client.post('/integrations/dify/retrieval',headers={'Authorization':'Bearer dify-test'},json=dict(body,knowledge_id='other')).status_code,404)
        self.assertEqual(self.client.post('/integrations/dify/retrieval',headers={'Authorization':'Bearer dify-test'},json=dict(body,metadata_condition={'conditions':[]})).status_code,422)

    def test_platform_auth_is_separate_and_does_not_accept_identity_fields(self):
        for platform in ('coze','dify'):
            path=f'/integrations/{platform}/ask'
            self.assertEqual(self.client.post(path,headers=self.headers(),json={'query':'退款审批'}).status_code,401)
            headers={'Authorization':'Bearer '+platform+'-test'}
            self.assertEqual(self.client.post(path,headers=headers,json={'query':'退款审批','tenant_id':'other'}).status_code,422)
            answer=self.client.post(path,headers=headers,json={'query':'发布窗口 密钥轮换'}).json()
            self.assertTrue(answer['refused'])

    def test_delegated_platform_requires_user_token_and_same_tenant(self):
        headers={'Authorization':'Bearer delegated-test'}
        self.assertEqual(self.client.post('/integrations/coze/ask',headers=headers,json={'query':'发布窗口'}).status_code,401)
        response=self.client.post('/integrations/coze/ask',headers={**headers,'X-RAG-User-Token':'bob-test'},json={'query':'发布窗口'})
        self.assertEqual(response.status_code,200,response.text)
        self.assertTrue(any(c['document_id'].startswith('E') for c in response.json()['citations']))
        self.assertEqual(self.client.post('/integrations/coze/ask',headers={**headers,'X-RAG-User-Token':'other-test'},json={'query':'发布窗口'}).status_code,403)

    def test_external_retrieval_does_not_invoke_answer_generator(self):
        with patch.object(self.app.state.service.generator,'generate',side_effect=AssertionError('should not generate')):
            response=self.client.post('/integrations/dify/retrieval',headers={'Authorization':'Bearer dify-test'},json={
                'knowledge_id':'enterprise','query':'差旅报销','retrieval_setting':{'top_k':1,'score_threshold':0.0}})
        self.assertEqual(response.status_code,200,response.text)


class ParsingTests(unittest.TestCase):
    def test_docx_paragraphs_tables_and_bad_files(self):
        from docx import Document
        document=Document();document.add_paragraph('A policy paragraph.');table=document.add_table(rows=1,cols=2)
        table.cell(0,0).text='Deadline';table.cell(0,1).text='12 days'
        stream=io.BytesIO();document.save(stream)
        text,locations=parse_upload('../policy.docx',stream.getvalue())
        self.assertIn('12 days',text);self.assertTrue(locations)
        with self.assertRaises(AccessError):parse_upload('bad.docx',b'not a zip')

    def test_docx_expansion_is_bounded(self):
        stream=io.BytesIO()
        with zipfile.ZipFile(stream,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('big.xml',b'x'*(41*1024*1024))
        with self.assertRaises(AccessError) as caught:parse_upload('big.docx',stream.getvalue())
        self.assertEqual(caught.exception.status,413)


class EvidenceAndContractsTests(unittest.TestCase):
    def test_failed_final_evidence_check_cannot_be_bypassed_by_resume(self):
        from rag_app.checkpoint import Checkpointer
        from rag_app.workflow import execute
        class Judge:
            model='fixture'
            final_calls=0
            def assess(self,query,chunks,proposed_answer=None):
                if proposed_answer is None:
                    return {'sufficient':True,'supporting_chunk_ids':[r['chunk_id'] for r in chunks]}
                self.final_calls+=1
                if self.final_calls==1:raise AccessError('evidence_judge_unavailable',503)
                return {'sufficient':False,'supporting_chunk_ids':[]}
        judge=Judge()
        with tempfile.TemporaryDirectory() as folder:
            store=Store(Path(folder)/'db');populate(store)
            service=Service(store,evidence_policy=EvidencePolicy(judge=judge))
            actor=Actor('alice','demo',('support',));cp=Checkpointer(Path(folder)/'cp')
            with self.assertRaises(AccessError):execute(service,cp,'退款审批',actor,workflow_id='same')
            result=execute(service,cp,'退款审批',actor,workflow_id='same')
            self.assertTrue(result['refused'])
            self.assertEqual(result['refusal_reason'],'synthesis_not_supported')
            self.assertEqual(judge.final_calls,2)

    def test_weak_match_missing_requested_detail_and_empty_query_refuse(self):
        policy=EvidencePolicy()
        rows=[{'chunk_id':'a','title':'入职材料','text':'事实：入职材料包括签收表。'}]
        self.assertEqual(policy.assess('工单升级的规定是什么？',rows)[0],[])
        rows=[{'chunk_id':'a','title':'退款审批','text':'事实：退款审批需要主管确认。'}]
        self.assertEqual(policy.assess('退款审批金额上限是多少？',rows)[1]['status'],'insufficient')
        self.assertEqual(policy.assess('',rows)[0],[])

    def test_model_judge_rejects_unknown_ids_and_failure_has_no_fallback(self):
        budget=Mock();judge=HTTPEvidenceJudge('https://invalid','secret','fixture',budget)
        response=httpx.Response(200,json={'choices':[{'message':{'content':'{"sufficient":true,"supporting_chunk_ids":["stolen"]}'}}]},request=httpx.Request('POST','https://invalid'))
        with patch('rag_app.evidence_policy.httpx.post',return_value=response):
            with self.assertRaises(AccessError):judge.assess('q',[{'chunk_id':'one','title':'t','text':'q'}])
        budget.assert_called_once_with('evidence')

    def test_model_rejection_refuses_without_generation(self):
        judge=Mock();judge.model='fixture';judge.assess.return_value={'sufficient':False,'supporting_chunk_ids':[]}
        rows=[{'chunk_id':'one','title':'退款审批','text':'退款审批上限500元。'}]
        selected,check=EvidencePolicy(judge=judge).assess('退款审批',rows)
        self.assertEqual(selected,[]);self.assertEqual(check['status'],'insufficient')

    def test_exported_tools_never_expose_document_writes_or_credentials(self):
        for platform in ('coze','dify'):
            spec=specification(platform,'https://rag.example.invalid')
            self.assertEqual(spec['openapi'],'3.0.3')
            self.assertFalse(any('/documents' in p for p in spec['paths']))
            for path in spec['paths'].values():
                props=path['post']['requestBody']['content']['application/json']['schema']['properties']
                self.assertFalse({'actor','token','groups','tenant_id'}&set(props))

    def test_missing_live_configuration_is_not_reported_as_verified(self):
        with patch.dict(os.environ,{},clear=True):
            report=run_probes(['postgres','embedding','rerank','model','langfuse','dify','coze'])
        self.assertEqual(report['summary']['verified'],0)
        self.assertEqual(report['summary']['not_configured'],7)


class IncrementalIndexTests(unittest.TestCase):
    def test_only_new_chunk_is_embedded_and_existing_vectors_are_kept(self):
        from rag_app.pg_native_index import NativePGIndex
        connection=Mock()
        def execute(sql,args=()):
            result=Mock();result.rowcount=0
            if 'SELECT chunk_id FROM embeddings' in sql:result.fetchall.return_value=[{'chunk_id':'old'}]
            elif "SELECT value FROM metadata" in sql:result.fetchone.return_value={'value':'2'}
            return result
        connection.execute.side_effect=execute
        store=Mock();store.connect.return_value.__enter__=Mock(return_value=connection)
        store.connect.return_value.__exit__=Mock(return_value=False)
        store.snapshot_all.return_value=('2',[{'chunk_id':'old','title':'old','text':'unchanged'},
                                             {'chunk_id':'new','title':'new','text':'changed'}])
        index=object.__new__(NativePGIndex);index.store=store;index.model_id='fixture'
        index.vectors=Mock(return_value=[[1.,0.]])
        self.assertEqual(index.sync(),1)
        index.vectors.assert_called_once_with(['new changed'])
        self.assertEqual(index.last_sync['retained_chunks'],1)
